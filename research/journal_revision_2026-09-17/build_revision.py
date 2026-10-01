from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
FIG=ROOT/'figures'
content=json.loads((ROOT/'manuscript_content.json').read_text())
TITLE=content['title'];AUTHORS=content['authors']
AFFIL='[Department and institution, city, country]'
EMAIL='[Corresponding author email]'
ABSTRACT=content['abstract'];KEYWORDS=content['keywords'];blocks=content['blocks']

doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75)
sec.left_margin=Inches(.625);sec.right_margin=Inches(.625)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
styles=doc.styles
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    styles[sty].font.name='Times New Roman';styles[sty].font.color.rgb=RGBColor(0,0,0)
    rpr=styles[sty].element.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for key in list(fonts.attrib):
            if 'Theme' in key:del fonts.attrib[key]
    ppr=styles[sty].element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
normal=styles['Normal'];normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=Pt(11.5);normal.paragraph_format.space_after=Pt(2)
normal.paragraph_format.first_line_indent=Inches(.12)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for sty in ['Heading 1','Heading 2']:
    styles[sty].font.size=Pt(10);styles[sty].paragraph_format.first_line_indent=Inches(0)
    styles[sty].paragraph_format.space_before=Pt(8);styles[sty].paragraph_format.space_after=Pt(4)
    styles[sty].paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Heading 1'].font.bold=False
styles['Heading 2'].font.bold=False;styles['Heading 2'].font.italic=True
styles['Caption'].font.size=Pt(8);styles['Caption'].font.italic=False
styles['Caption'].font.bold=False
styles['Caption'].paragraph_format.line_spacing=Pt(9)
styles['Caption'].paragraph_format.first_line_indent=Inches(0)
styles['Caption'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(22);styles['Title'].font.bold=False
styles['Title'].paragraph_format.line_spacing=Pt(25)
styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Title'].paragraph_format.first_line_indent=Inches(0)
p=doc.add_paragraph(TITLE,'Title');p.paragraph_format.space_after=Pt(9)
for val in [AUTHORS,AFFIL,EMAIL]:
    p=doc.add_paragraph(val);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(2)
    for run in p.runs:run.font.size=Pt(10)
sec=doc.add_section(WD_SECTION_START.CONTINUOUS)
cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360')
p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0)
p.add_run('Abstract - ').bold=True;p.add_run(ABSTRACT)
for r in p.runs:r.font.size=Pt(9);r.bold=True
p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0)
p.add_run('Index Terms - ').bold=True;p.add_run(KEYWORDS)
for r in p.runs:r.font.size=Pt(9)

def mr(value):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r
def math_container(tag,children):
    node=OxmlElement('m:'+tag)
    for x in children:node.append(x)
    return node
def frac(a,b):return math_container('f',[math_container('num',a),math_container('den',b)])
def sub(a,b):return math_container('sSub',[math_container('e',[mr(a)]),math_container('sub',[mr(b)])])
def sup(children,power):return math_container('sSup',[math_container('e',children),math_container('sup',[mr(power)])])
def radical(children):
    prop=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');prop.append(hide)
    return math_container('rad',[prop,math_container('deg',[]),math_container('e',children)])
def mathline(nodes):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
    p._p.append(math_container('oMath',nodes))
    return p

for b in blocks:
    typ=b['type']
    if typ in ['p','h1','h2','ref']:
        p=doc.add_paragraph(b['text'],{'h1':'Heading 1','h2':'Heading 2'}.get(typ,'Normal'))
        if typ=='ref':
            p.paragraph_format.first_line_indent=Inches(-.18);p.paragraph_format.left_indent=Inches(.18)
            p.paragraph_format.line_spacing=Pt(9);p.paragraph_format.space_after=Pt(4)
            for r in p.runs:r.font.size=Pt(8)
    elif typ=='eq':
        mu=chr(956);sigma=chr(931);pi=chr(960);phi=chr(966)
        if '(1)' in b['text']:
            mathline([mr(mu+' = '),frac([mr('1')],[mr('n')]),mr(sigma+' '),sub('x','i')]).paragraph_format.keep_with_next=True
            mathline([mr('s = '),radical([frac([mr(sigma+' '),sup([mr('('),sub('x','i'),mr(' - '+mu+')')],'2')],[mr('n - 1')])]),mr('   (1)')])
        elif '(2)' in b['text']:
            mathline([mr('z = '),frac([sub('x','current'),mr(' - '+mu)],[mr('max(s, '),sub('s','min'),mr(')')]),mr('   (2)')])
        elif '(3)' in b['text']:
            mathline([sub('R','0'),mr(' = 0.18S + 0.22C + 0.18E')]).paragraph_format.keep_with_next=True
            mathline([mr('+ 0.26F + 0.10O + 0.06T   (3)')])
        elif '(5)' in b['text']:
            mathline([mr('d = w '),frac([mr('A - '+chr(952))],[mr('A - 1')]),mr('   (5)')])
        elif '(6)' in b['text']:
            mathline([sub('P','capture'),mr('(f) = min(1, fd)   (6)')])
        else:
            mathline([mr('a(t) = 1 + 0.004 sin(2'+pi+' 0.7t + '+phi+')')]).paragraph_format.keep_with_next=True
            mathline([mr('+ (A - 1) max(0, 1 - '),frac([mr('2|t - '),sub('t','0'),mr('|')],[mr('w')]),mr(')   (4)')])
    elif typ=='fig':
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(FIG/(b['name']+'.png')),width=Inches(3.48))
        doc.add_paragraph(b['caption'],'Caption')
    elif typ=='table':
        p=doc.add_paragraph(b['caption'],'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        t=doc.add_table(rows=1,cols=len(b['head']));t.autofit=False
        for cell,width,title in zip(t.rows[0].cells,b['widths'],b['head']):cell.width=Inches(width);cell.text=title
        for row in b['rows']:
            cells=t.add_row().cells
            for cell,width,val in zip(cells,b['widths'],row):cell.width=Inches(width);cell.text=str(val)
        for col,width in zip(t.columns,b['widths']):col.width=Inches(width)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            tag=OxmlElement('w:'+side);tag.set(qn('w:val'),'single');tag.set(qn('w:sz'),'4');tag.set(qn('w:color'),'D9D9D9');borders.append(tag)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),'55');e.set(qn('w:type'),'dxa');marg.append(e)
                cp.append(marg)
                if i==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EDEFF1');cp.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(9)
                    p.paragraph_format.keep_with_next=(i<len(t.rows)-1)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:r.font.size=Pt(8);r.bold=(i==0)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)

for section in doc.sections[:1]:
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    foot.paragraph_format.first_line_indent=Inches(0)

# Remove inherited decorative paragraph borders and theme font overrides.
for style in styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for fonts in style.element.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if 'Theme' in key:del fonts.attrib[key]
doc.core_properties.title=TITLE;doc.core_properties.author='';doc.core_properties.subject='Synthetic software evaluation of NeuroGuardian X'
doc.core_properties.comments='Editable IEEE-style journal manuscript. Author details require completion.'
doc.save(ROOT/'NeuroGuardian_X_IEEE_Journal_Revision.docx')

def texesc(s):
    for a,b in [('\\',r'\textbackslash{}'),('&',r'\&'),('%',r'\%'),('$',r'\$'),('#',r'\#'),('_',r'\_')]:s=s.replace(a,b)
    return s
tex=[r'\documentclass[journal]{IEEEtran}',r'\usepackage{graphicx,amsmath,array,url}',
     r'\graphicspath{{figures/}}',r'\begin{document}',r'\title{'+texesc(TITLE)+'}',
     r'\author{'+texesc(AUTHORS)+r'\thanks{'+texesc(AFFIL+'; '+EMAIL)+r'}}',r'\maketitle',
     r'\begin{abstract}'+texesc(ABSTRACT)+r'\end{abstract}',r'\begin{IEEEkeywords}'+texesc(KEYWORDS)+r'\end{IEEEkeywords}']
inrefs=False
for b in blocks:
    typ=b['type']
    if typ=='h1':
        title=b['text']
        if title=='REFERENCES':tex.append(r'\begin{thebibliography}{11}');inrefs=True
        elif title[0] in 'IVX' and '. ' in title:tex.append(r'\section{'+texesc(title.split('. ',1)[1].title())+'}')
        else:tex.append(r'\section*{'+texesc(title.title())+'}')
    elif typ=='h2':tex.append(r'\subsection{'+texesc(b['text'].split('. ',1)[1])+'}')
    elif typ=='p':tex.append(texesc(b['text'])+'\n')
    elif typ=='eq':tex.append(r'\begin{equation}'+b['latex']+r'\end{equation}')
    elif typ=='fig':
        cap=b['caption'].split('. ',2)[2]
        tex.append(r'\begin{figure}[!t]\centering\includegraphics[width=\columnwidth]{'+b['name']+r'.pdf}\caption{'+texesc(cap)+r'}\end{figure}')
    elif typ=='table':
        spec=''.join('p{'+f'{w/3.5*.91:.3f}'+r'\columnwidth}' for w in b['widths'])
        cap=b['caption'].split('. ',1)[1]
        tex.extend([r'\begin{table}[!t]\caption{'+texesc(cap)+r'}\centering\scriptsize\setlength{\tabcolsep}{2pt}',r'\begin{tabular}{'+spec+r'}\hline',
                    ' & '.join(texesc(x) for x in b['head'])+r'\\\hline'])
        tex += [' & '.join(texesc(str(x)) for x in row)+r'\\' for row in b['rows']]
        tex.append(r'\hline\end{tabular}\end{table}')
    elif typ=='ref':
        n=b['text'].split(']')[0][1:]
        raw=b['text'].split('] ',1)[1]
        chunks=re.split(r'(https?://\S+)',raw)
        rendered=''.join(r'\url{'+x+'}' if x.startswith('http') else texesc(x) for x in chunks)
        tex.append(r'\bibitem{ref'+n+'}'+rendered)
if inrefs:tex.append(r'\end{thebibliography}')
tex.append(r'\end{document}')
(ROOT/'NeuroGuardian_X_IEEE_Journal_Revision.tex').write_text('\n'.join(tex),encoding='utf-8')
