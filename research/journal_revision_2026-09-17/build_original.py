from pathlib import Path
import csv
import hashlib
import json
import math
import re
import shutil
import statistics as st
import subprocess
import sys
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib import colors

ROOT=Path(__file__).resolve().parent
APP=ROOT.parents[1]/'neuroguardian_app'
RESULTS=ROOT/'results'
FIG=ROOT/'figures';FIG.mkdir(exist_ok=True)
rows=json.loads((RESULTS/'neurotwin_trials.json').read_text())
motion=json.loads((RESULTS/'motion_summary.json').read_text())
control=json.loads((RESULTS/'controller_cases.json').read_text())
summary={}
for s in dict.fromkeys(r['scenario'] for r in rows):
    a=[r for r in rows if r['scenario']==s]
    summary[s]={'n':len(a),'risk_mean':st.mean(r['risk_score'] for r in a),
        'risk_sd':st.stdev(r['risk_score'] for r in a),
        'stress_mean':st.mean(r['stress_score'] for r in a),
        'ecg_mean':st.mean(r['ecg_score'] for r in a),
        'quality_mean':st.mean(r['quality'] for r in a),
        'baseline_ready':sum(r['baseline_ready'] for r in a),
        'quality_reliable':sum(r['quality_reliable'] for r in a)}
(RESULTS/'neurotwin_summary.json').write_text(json.dumps(summary,indent=2))
with (RESULTS/'neurotwin_trials.csv').open('w',newline='') as f:
    fields=[k for k in rows[0] if k not in ['reasons','deviations']]
    w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(rows)

BLUE=colors.HexColor('#244764');RED=colors.HexColor('#a43b35');GRAY=colors.HexColor('#70777d')
def fig_canvas(name,w=252,h=220):
    c=canvas.Canvas(str(FIG/(name+'.pdf')),pagesize=(w,h))
    c.setTitle(name.replace('_',' '));return c
def text(c,x,y,t,size=8,font='Times-Roman',align='left'):
    c.setFillColor(colors.black);c.setFont(font,size)
    (c.drawCentredString if align=='center' else c.drawRightString if align=='right' else c.drawString)(x,y,t)
def arrow(c,x1,y1,x2,y2):
    c.setStrokeColor(colors.black);c.setLineWidth(.6);c.line(x1,y1,x2,y2)
    a=math.atan2(y2-y1,x2-x1)
    for d in [-.5,.5]:c.line(x2,y2,x2-4*math.cos(a+d),y2-4*math.sin(a+d))

c=fig_canvas('architecture',252,268)
boxes=[(218,'ESP32-S3 sensor interface','IMU, optical, ECG, GSR, temperature'),
       (162,'BLE version 2 packets','Values, quality, movement and fall flags'),
       (106,'Flutter and NeuroTwin','Baseline, rule scores, 30-second controller'),
       (50,'FastAPI notification backend','Location, facility lookup, provider requests')]
for y,title,sub in boxes:
    c.setStrokeColor(GRAY);c.setFillColor(colors.HexColor('#f5f6f7'));c.rect(6,y,240,38,fill=1,stroke=1)
    text(c,126,y+24,title,9,'Times-Bold','center');text(c,126,y+10,sub,8,align='center')
for y in [218,162,106]:arrow(c,126,y,126,y-18)
arrow(c,126,50,126,31);text(c,126,18,'Guardian channel acceptance',9,'Times-Bold','center')
text(c,126,5,'Delivery and human acknowledgment are separate',7.6,align='center')
c.save()

c=fig_canvas('motion_results',252,235)
names=['Quiet','Walking','Impact + still','Recovery','HR drop','Unworn impact']
groups=['quiet','walking','impact_immobile','impact_recovery','impact_hr_drop','unworn_impact']
policies=['original_1Hz','original_50Hz','guarded_50Hz']
left=72;right=243;bottom=29;top=200
for p,color,label,x in zip(policies,[GRAY,BLUE,RED],['Original 1 Hz','Original 50 Hz','Guarded 50 Hz'],[2,87,176]):
    c.setFillColor(color);c.rect(x,221,6,6,fill=1,stroke=0);text(c,x+9,221,label,7)
for v in [0,25,50,75,100]:
    x=left+(right-left)*v/100;c.setStrokeColor(colors.HexColor('#dddddd'));c.line(x,bottom,x,top)
    text(c,x,17,str(v),7,align='center')
for j,(group,label) in enumerate(zip(groups,names)):
    y=top-j*28-10;text(c,left-5,y,label,7.4,align='right')
    for k,(p,color) in enumerate(zip(policies,[GRAY,BLUE,RED])):
        row=next(r for r in motion if r['group']==group and r['policy']==p)
        val=row['critical_percent'];yy=y+7-k*6
        c.setFillColor(color)
        if val>0:c.rect(left,yy,(right-left)*val/100,4,fill=1,stroke=0)
text(c,158,3,'Trials with critical flag (%)',8,align='center');c.save()

c=fig_canvas('risk_results',252,228)
labels=['Normal','Stationary','Stress shift','Exercise','ECG scalar','Poor contact','Missing','Fall inactive']
left=71;right=242;top=210;bottom=26
for v in [0,20,40,60,80,100]:
    x=left+(right-left)*v/100;c.setStrokeColor(colors.HexColor('#dddddd'));c.line(x,bottom,x,top)
    text(c,x,15,str(v),7,align='center')
for threshold in [30,60,80]:
    x=left+(right-left)*threshold/100;c.setDash(2,2);c.setStrokeColor(GRAY);c.line(x,bottom,x,top);c.setDash()
for j,(s,label) in enumerate(zip(summary,labels)):
    y=top-14-j*22;d=summary[s];value=d['risk_mean'];sd=d['risk_sd']
    text(c,left-5,y,label,7.5,align='right');c.setFillColor(BLUE)
    c.rect(left,y-2,(right-left)*value/100,9,fill=1,stroke=0)
    x=left+(right-left)*value/100;err=(right-left)*sd/100
    c.setStrokeColor(colors.black);c.line(x-err,y+2.5,x+err,y+2.5)
    text(c,min(right-14,x+5),y,f'{value:.1f}',7)
text(c,157,2,'Risk score, mean +/- sample SD',8,align='center');c.save()

c=fig_canvas('sampling_failure',252,226)
left=33;right=243;top=211;bottom=124
def px(t):return left+(t-5.7)/1.6*(right-left)
def py(g):return bottom+(g/4)*(top-bottom)
for g in [0,1,2,3,4]:
    text(c,27,py(g)-2,str(g),7,align='right');c.setStrokeColor(colors.HexColor('#dddddd'));c.line(left,py(g),right,py(g))
c.setStrokeColor(BLUE);c.setLineWidth(1)
points=[]
for i in range(321):
    t=5.7+i*.005;g=1+max(0,1-abs(t-5.98)/.1)*2.4;points.append((px(t),py(g)))
p=c.beginPath();p.moveTo(*points[0])
for q in points[1:]:p.lineTo(*q)
c.drawPath(p)
c.setDash(2,2);c.setStrokeColor(GRAY);c.line(left,py(2.6),right,py(2.6));c.setDash()
text(c,174,py(2.6)+3,'2.6 g threshold',7)
for t in [6,7]:
    g=1+max(0,1-abs(t-5.98)/.1)*2.4;c.setFillColor(RED);c.circle(px(t),py(g),2,fill=1,stroke=0)
    text(c,px(t),113,f'{t:.0f}',7,align='center')
text(c,138,100,'Time (s); dots are 1 Hz samples',8,align='center');text(c,4,215,'g',8)
text(c,8,79,'Original: impact at 6 s, cancellation at 7 s',8)
text(c,8,65,'The decaying impact creates a large change in g.',8)
text(c,8,44,'Guarded candidate: preserve the countdown',8,'Times-Bold')
text(c,8,30,'Ignore the initial settling transient; require',8)
text(c,8,17,'2 s of sustained movement evidence to cancel.',8)
text(c,8,3,'Illustrative trace; not a measured human fall.',7)
c.save()

poppler=Path(r'C:\Users\hp\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe')
for p in FIG.glob('*.pdf'):
    subprocess.run([str(poppler),'-png','-r','300','-singlefile',str(p),str(p.with_suffix(''))],check=True,capture_output=True)

blocks=[]
def para(t):blocks.append({'type':'p','text':t})
def h1(t):blocks.append({'type':'h1','text':t})
def h2(t):blocks.append({'type':'h2','text':t})
def eq(t,latex):blocks.append({'type':'eq','text':t,'latex':latex})
def figure(name,caption):blocks.append({'type':'fig','name':name,'caption':caption})
def table(caption,head,rows,widths):blocks.append({'type':'table','caption':caption,'head':head,'rows':rows,'widths':widths})

TITLE='NeuroGuardian X Reproducible Simulation and Failure Analysis of Personalized Wearable Monitoring and Fall Alert Logic'
AUTHORS='[Author names in publication order]'
AFFIL='[Department and institution, city, country]'
EMAIL='[Corresponding author email]'
ABSTRACT=('Wearable safety systems combine physiological sensing, motion processing, smartphone software, and remote notification services. '
'A plausible dashboard does not establish that this chain behaves correctly during a fall or communication failure. This paper presents a reproducible software evaluation of NeuroGuardian X, '
'an ESP32-S3 and Flutter prototype with a personal statistical profile called NeuroTwin and a FastAPI notification backend. The study executes the existing Flutter engine on 1,600 synthetic cases, '
'replays 1,200 generated motion traces through three policies, and exercises five mobile-controller cases and five backend contract cases. A verbatim extraction of the firmware motion function '
'captures 14 of 200 short impact-and-immobility traces at 1 Hz, but produces no critical flag because post-impact changes cancel the timer. Raising sampling to 50 Hz increases capture to 199 of 200 '
'without correcting that cancellation. An experimental guarded recovery policy produces critical flags in 199 of 200 immobile-impact traces, but also in 198 of 200 unworn-device impact traces. '
'The Flutter engine labels all 200 missing-signal cases Normal, and its baseline rejects all 200 stationary histories. Backend results distinguish request acceptance from unverified delivery. '
'These findings identify integration defects and a reproducible path for their investigation. They do not establish clinical accuracy, real-world fall sensitivity, or reliable emergency dispatch.')
KEYWORDS='Wearable sensing, fall detection, software-in-the-loop, signal quality, personalized baseline, emergency notification, reproducibility.'

h1('I. INTRODUCTION')
para('A wearable safety application must do more than identify a large sensor value. It must preserve an event across transient motion, distinguish unavailable measurements from ordinary measurements, '
'maintain a meaningful personal reference, and communicate what has actually happened to an alert. Errors at these interfaces can remain hidden when development focuses on dashboard appearance or isolated unit tests. '
'The engineering question addressed here is whether an implemented prototype behaves consistently with its stated monitoring and escalation rules under controlled inputs.')
para('NeuroGuardian X combines an ESP32-S3 sensor interface, a Bluetooth Low Energy (BLE) connection, a Flutter mobile application, and a Python notification backend. NeuroTwin is the project name for a '
'personal statistical baseline and rule-based explanation engine. In this implementation it is not a mechanistic physiological model, a disease simulator, or a clinically validated digital twin. '
'The current study evaluates software behavior before making claims about measurements obtained from a completed wrist-worn device.')
para('The contribution is a traceable evaluation across three implementation boundaries. First, a host harness executes the firmware motion-state function without rewriting its decision branches. '
'Second, synthetic histories are passed to the production Flutter engine and controller. Third, backend endpoint functions are exercised with controlled provider outcomes. '
'An experimental recovery policy is evaluated as a diagnostic comparison, including an intentionally difficult unworn-device case. The work reports negative results explicitly and provides the inputs, '
'generators, source hashes, and numeric outputs needed to reproduce them.')
para('The scope is software-in-the-loop testing. Sensor electronics, battery life, skin contact, over-the-air BLE timing, Android background execution, internet transport, and actual guardian response are outside the measured system. '
'No person was asked to fall, no patient recording was processed, and no emergency message was sent during these experiments. All use of the words normal, stress, exercise, and fall in scenario names refers to constructed inputs.')

h1('II. RELATED WORK AND EVALUATION SCOPE')
para('SisFall supplies recorded falls and activities of daily living and demonstrates why performance must be examined by activity and population [1]. Its original recordings are sampled at 200 Hz. '
'The present synthetic waveforms are neither taken from SisFall nor a substitute for its recordings. Their short impulses are deliberately chosen to probe sampling and cancellation behavior. '
'Thus, a critical-flag percentage reported here cannot be compared directly with a published fall-classification accuracy.')
para('Hongn et al. provide wrist-worn physiological recordings under acute stress and structured exercise [2]. Their work motivates evaluating activity context when interpreting pulse and electrodermal measurements. '
'However, plausible combinations of increased heart rate and electrodermal activity do not create clinically labeled stress examples. In this paper, those combinations probe the software response to a feature shift; '
'there is no trained stress classifier and no measured stress-detection accuracy.')
para('The MIT-BIH Arrhythmia Database provides 48 half-hour two-channel ECG records sampled at 360 Hz [3], [4]. It is relevant to a future waveform-based ECG evaluation, whereas the current prototype sends a scalar ECG value '
'in its periodic packet. A scalar amplitude deviation cannot be equated with rhythm classification. The exam-stress dataset [5], Sudden Cardiac Death Holter Database [6], and induced-stress/exercise dataset release [7] '
'are also relevant future resources. None of their records was downloaded, used for fitting, or used for testing in this study. The amount of those datasets used for training is zero.')
para('Our research questions are: RQ1, how do the implemented baseline, quality, and score rules respond to missing, stationary, and shifted inputs; RQ2, how do sampling frequency and recovery cancellation interact; '
'and RQ3, what evidence is actually available after a controller trigger or backend success response? The comparison is internal to the prototype, rather than a claim of superiority to commercial watches.')

h1('III. SYSTEM ARCHITECTURE AND IMPLEMENTATION')
figure('architecture','Fig. 1. Prototype data path. Software boundaries are exercised independently. Arrows do not imply that the complete hardware-to-guardian path was tested.')
h2('A. Wearable and packet interface')
para('The inspected firmware uses Arduino C++ on an ESP32-S3. It contains an MPU6050 motion path, analog GSR and ECG inputs, a battery input, an SOS button, and buzzer/vibration control. '
'Optional compile-time paths exist for a MAX30102 optical module and a TinyGPSPlus GNSS receiver; both are disabled in the inspected configuration. Pressure and altitude are placeholders, and a dedicated calibrated '
'temperature measurement path is unfinished. The MAX30102 internal temperature path, when enabled, should not be interpreted as an independently validated body-temperature measurement.')
para('Version 2 packets carry compact JSON fields for timestamp, heart rate, oxygen saturation, HRV, GSR, temperature, ECG scalar, motion axes, quality scores, position, battery, and event flags. '
'The BLE maximum transmission unit is requested as 517 bytes, but successful negotiation and packet delivery are not measured here. The firmware calls notifyMetrics at a nominal one-second interval only while '
'deviceConnected is true. Motion acquisition and buildMotionState are inside that function. Consequently, the configured source ties motion-state advancement to the connected notification loop; the nominal rate is an '
'idealization that excludes scheduling jitter and missed intervals.')
h2('B. Mobile application and NeuroTwin')
para('The Flutter application parses the packet into a Metrics object. NeuroTwin combines a signal-quality view, an activity-dependent personal baseline, trend summaries, and weighted rule scores. '
'The baseline uses at most 120 eligible preceding samples. An overall profile is marked ready at 20 eligible samples; individual metric readiness uses 12 samples. These sample-count checks are software conventions '
'and do not establish that a stable multi-day physiological baseline has been learned.')
para('For metric x with n eligible samples, the baseline mean, sample standard deviation, and current normalized deviation are computed as follows. A metric-specific floor prevents division by a very small standard deviation.')
eq('mu = (1/n) sum(x_i); s = sqrt(sum((x_i - mu)^2)/(n - 1))                 (1)',r'\begin{aligned}\mu&=\frac{1}{n}\sum_{i=1}^{n}x_i,\\s&=\sqrt{\frac{\sum_{i=1}^{n}(x_i-\mu)^2}{n-1}}.\end{aligned}')
eq('z = (x_current - mu) / max(s, s_min)                                  (2)',r'z=\frac{x_{\mathrm{current}}-\mu}{\max(s,s_{\min})}.')
para('The standard-deviation floors are 3 beats/min for heart rate, 0.8 percentage points for oxygen saturation, 5 ms for HRV, 4 arbitrary units for GSR, 0.2 degrees C for temperature, and 0.08 mV for the ECG scalar. '
'Eligibility requires acceptable contact, an overall quality of at least 0.45, compatible activity, and some usable physiological input. It also rejects every sample with movementDetected equal to false. '
'That last condition is evaluated directly in the stationary-history experiment.')
para('Quality values are provided by the firmware or estimated from value availability and motion. NeuroTwin averages PPG, ECG, GSR, temperature, and motion quality and declares the result reliable when the average is at '
'least 0.55 and watch fit is acceptable. These quality numbers are heuristics, not empirically calibrated probabilities. GPS readiness is handled separately. A high aggregate quality does not validate every individual channel.')
h2('C. Composite score and labels')
para('Let S denote the stress score, C the existing cardiac-related rule score, E the ECG-scalar anomaly score, F the motion/fall score, O the oxygen-related load, and T the temperature-related load, each on a 0-100 scale. '
'The current composite is:')
eq('R0 = 0.18S + 0.22C + 0.18E + 0.26F + 0.10O + 0.06T                  (3)',r'R_0=0.18S+0.22C+0.18E+0.26F+0.10O+0.06T.')
para('For unreliable quality without a fall, the score is multiplied by 0.72. A fall with no movement sets a minimum score of 82, and the Critical fall activity sets a minimum of 92. '
'A non-fall case without a ready baseline is capped at 72. Labels change at scores of 30, 60, and 80: Normal, Observe, Warning, and Critical Review. Although the application uses probability terminology for some '
'components, this paper calls them scores because no probability calibration has been performed.')
para('The stress calculation combines deviations in GSR, HRV, heart rate, and temperature and blends them with recent transmitted stress values. Exercise activity changes that calculation. The ECG component combines '
'scalar-amplitude deviation, baseline deviation, and HRV load. These operations explain why isolated component scores may rise without changing the final label. Trend windows of five minutes, thirty minutes, and one hour '
'are implemented; the short histories in this experiment do not validate those time scales.')
h2('D. Alert controller and backend')
para('The mobile controller starts a 30-second countdown following a fall signal. A subsequent non-fall movement packet cancels the countdown. An immediate critical path uses a heart-rate drop of at least 20 beats/min, '
'whereas the firmware constant is 22 beats/min. A Critical fall activity also triggers the mobile critical branch. The timer uses DateTime.now and periodic callbacks; its behavior is therefore tested with actual elapsed '
'time rather than assuming that a test framework virtual clock changes DateTime.now.')
para('The backend accepts an alert containing event information, coordinates, guardian settings, and optional medical facilities. If the list is absent, it can query Google Places Nearby Search [8]. '
'Facility discovery returns information about places; it is not a hospital dispatch agreement. WhatsApp and push are separate provider paths. The current response includes success and delivery fields based on '
'provider-call outcomes, while the webhook handler acknowledges receipt without retaining a message-delivery history. The controlled backend tests examine this contract, not a live service.')

h1('IV. EXPERIMENTAL METHODS')
h2('A. Provenance and execution boundaries')
para('Experiments were executed on Windows on 17 September 2026 against application version 0.4.0+4, using Flutter 3.19.5 and Dart 3.3.3. The backend environment used Python 3.10.10, FastAPI 0.141.1, and HTTPX 0.28.1. '
'The native harness was compiled with the installed Microsoft C++ toolchain, version family 14.50. Production source files were read but not modified. A manifest records SHA-256 hashes of the inspected files; '
'the working tree contains changes beyond its Git commit, so the file hashes define the evaluated state more precisely than the commit alone.')
para('For motion replay, the harness extracts the complete buildMotionState function, its MotionSample structure, and its relevant constants verbatim from the firmware. The host substitutes a controllable millisecond clock, '
'a released SOS button, and inert output functions. The extracted function runs as native C++ with single-precision motion values. This is direct execution of an isolated decision function, not emulation of an ESP32 '
'processor, Arduino scheduler, sensor bus, or BLE stack. Connected operation is assumed in the motion trials.')
para('The NeuroTwin experiments import the production Dart classes through the Flutter test runner. The backend experiments import the production endpoint functions, disable dotenv loading, remove provider configuration '
'from the experiment environment, and replace notification and facility providers with local stubs. Network-client construction is blocked. This setup prevents a synthetic event from reaching a real guardian or hospital.')
h2('B. NeuroTwin synthetic cases')
para('Five generator seeds, 17, 29, 43, 71, and 101, each produce 40 repetitions per scenario, giving 200 cases per scenario and 1,600 engine evaluations. Every case contains 120 preceding one-second observations and '
'one current observation. The current observation is excluded from its baseline. Except in the stationary scenario, preceding observations represent the ordinary resting input range with movementDetected set to true. '
'This flag choice is intentional: it permits a comparison with stationary resting histories, rather than asserting that those histories were measured from people.')
table('TABLE I. CURRENT SYNTHETIC INPUT RANGES',['Scenario','Changed inputs'],[
    ['Normal','HR 68-76; SpO2 97-99; HRV 45-55; GSR 28-36'],
    ['Stationary','Same ranges; all movement flags false'],
    ['Stress shift','HR 95-115; HRV 15-25; GSR 75-90; stress 75-90'],
    ['Exercise','HR 120-145; HRV 20-35; activity Running'],
    ['ECG scalar shift','ECG 1.6-2.0 mV; other values ordinary'],
    ['Poor contact','Quality 0.05-0.20; fit false; shifted numeric values'],
    ['Missing','HR and temperature zero; optional vitals absent'],
    ['Fall inactive','Fall flag true; movement false; ordinary vitals']], [1.0,2.5])
para('Ordinary temperature is sampled uniformly from 36.5 to 36.8 degrees C, ECG scalar from 0.55 to 0.69 mV, transmitted stress from 25 to 35, and supplied channel quality from 0.85 to 0.95. '
'Integer ranges include their stated endpoints; continuous upper limits are excluded by the generator. GSR is an arbitrary 0-100 value rather than calibrated conductance. HRV is an input scalar, not calculated from '
'synthetic beat waveforms. These ranges are engineering test choices and are not normative clinical intervals.')
para('The poor-contact case keeps plausible but strongly shifted numbers while setting contact quality low. This separates channel presence from channel trust. The missing case supplies no physiological measurements '
'but preserves an ordinary motion vector. The ECG-scalar case is intentionally not called an arrhythmia or cardiac-arrest case. For each evaluation, the harness stores component scores, final score and category, '
'quality reliability, baseline eligibility, explanations, and input values. Results are summarized using the mean and sample standard deviation, without treating scenario labels as disease ground truth.')
h2('C. Motion generator and sampling comparison')
para('The motion experiment uses the same five seeds and 40 repetitions per seed for each of six trace families: quiet, walking, impact followed by immobility, impact followed by recovery, impact with a heart-rate step, '
'and an unworn-device impact. There are 1,200 parameterized traces. Each lasts 45 seconds and is independently replayed at 1 Hz and 50 Hz. For an impact trace, peak time t0 is uniform on [5,6) seconds, width w '
'is uniform on [0.08,0.24) seconds, and peak magnitude A is uniform on [2.8,4.2) g. A low-amplitude sinusoid is added to the 1 g baseline.')
eq('a(t) = 1 + 0.004 sin(2 pi 0.7t + phi) + (A - 1) max(0, 1 - 2|t-t0|/w)    (4)',r'\begin{aligned}a(t)&=1+0.004\sin(2\pi\,0.7t+\phi)\\&\quad+(A-1)\max\!\left(0,1-\frac{2|t-t_0|}{w}\right).\end{aligned}')
para('The phase phi is uniform on [0,2 pi). Quiet and walking traces omit the impact term. Walking adds a 0.35 g sinusoid at 1.8 Hz; recovery adds the same oscillation eight seconds after the peak. '
'The heart-rate-step trace changes from 75 to 45 beats/min at the start of the impact support. The unworn-device trace intentionally uses the same scalar motion family as the immobile-impact trace, because a magnitude '
'signal alone does not encode whether the device is attached to a person. The generator does not model orientation, skin contact, sensor saturation, or a population distribution of real falls.')
para('The production policy starts a countdown when magnitude exceeds 2.6 g. It defines movement as the absolute difference between successive magnitudes exceeding 0.10 g. If the timer is active and movement is '
'detected without a concurrent impact, it cancels the timer. A critical flag follows 30 seconds without cancellation or a qualifying concurrent heart-rate drop. Replay uses the same continuous trace parameters for '
'all policies; it does not optimize separate inputs to favor an alternative.')
figure('sampling_failure','Fig. 2. Illustrative synthetic counterexample. A sample on the impulse starts the timer, but the next sample near 1 g produces a large magnitude difference and cancels it.')
h2('D. Experimental guarded recovery policy')
para('A separate Python policy is evaluated only as an exploratory diagnostic alternative. It samples at 50 Hz, latches the same 2.6 g impact, and retains the same 30-second timeout. It ignores recovery cancellation '
'for the first 1.5 seconds. Thereafter, movement evidence is a peak-to-peak magnitude range greater than 0.18 g over the previous 0.5 seconds. Cancellation requires that evidence to persist for two seconds. '
'The candidate retains an adjacent-sample heart-rate-drop check, allowing its timing weakness to remain visible. It is not installed in the app or firmware.')
para('The guarded-policy values are engineering choices for exploring the observed cancellation defect. They were not learned from patient data or selected through a held-out validation procedure. The original policy '
'is replayed at 50 Hz as a sampling-only ablation, separating increased temporal resolution from the changed cancellation rule. Candidate and original code paths use different host languages; the principal endpoint '
'is a discrete state transition, not execution-time speed.')
h2('E. Endpoints and controller contracts')
para('For each policy and trace, the harness records whether a fall/countdown flag ever appears, whether the timer is canceled, and the first critical-flag time. This produces 3,600 policy-trace evaluations. '
'Capture and escalation are reported separately. A critical flag is a firmware decision, not a message sent to a guardian. Percentages use 200 trials per family and policy. Wilson 95% intervals quantify binomial '
'uncertainty under the chosen synthetic generator only; they do not quantify uncertainty over real users or environments.')
para('Five production-controller cases exercise fall followed by packet silence, movement after five seconds, manual reset, stationary input without a fall, and a Critical fall activity. The no-response cases run '
'concurrently for 32 real seconds; one observed timeout is reported rather than a latency distribution. Four backend cases cover every combination of accepted/rejected WhatsApp and push stubs, and a fifth injects a '
'failed-delivery webhook. No BLE disconnection is physically induced and no Android lock-screen or process-suspension behavior is tested.')

h1('V. RESULTS')
h2('A. NeuroTwin score and baseline behavior')
table('TABLE II. PRODUCTION ENGINE OUTPUTS FOR 200 CASES PER ROW',['Scenario','Risk mean (SD)','Ready / usable'],[
    [label,f"{summary[s]['risk_mean']:.2f} ({summary[s]['risk_sd']:.2f})",f"{summary[s]['baseline_ready']} / {summary[s]['quality_reliable']}"]
    for s,label in zip(summary,labels)], [1.1,1.18,1.22])
para('In Table II, Ready is the number of ready baseline profiles and usable is the number passing the current aggregate-quality rule. All normal histories become ready, whereas no stationary history does. '
'All exercise cases also have an unready baseline because their preceding history is resting and the current context is running. This latter result is expected for a new activity context; the stationary exclusion '
'is more consequential because quiet rest is a legitimate monitoring state. All counts refer to the program rules, not a validated quality reference.')
figure('risk_results','Fig. 3. Production NeuroTwin risk means with sample-standard-deviation error bars. Dashed lines mark category thresholds at 30, 60, and 80; none is a medically calibrated threshold.')
para(f"Normal and stationary cases have mean scores of {summary['normal']['risk_mean']:.2f} and {summary['stationary']['risk_mean']:.2f}, respectively. "
f"The {summary['stationary']['risk_mean']-summary['normal']['risk_mean']:.2f}-point difference follows the stillness contribution to the motion/fall score, although the scenario contains no fall. "
f"The stress-shift component rises to {summary['stress_shift']['stress_mean']:.2f} on average, but the composite remains {summary['stress_shift']['risk_mean']:.2f}. "
f"Similarly, the ECG-scalar component averages {summary['ecg_scalar_shift']['ecg_mean']:.2f}, while its composite averages {summary['ecg_scalar_shift']['risk_mean']:.2f}. "
'All 200 cases in each of these groups retain the Normal category. This is a score-aggregation observation, not evidence that the corresponding synthetic shifts are medically harmless.')
para(f"All 200 poor-contact inputs are marked unreliable, and their ECG component is zero, yet their stress component remains elevated at a mean of {summary['poor_contact']['stress_mean']:.2f}. "
'This occurs because the separate stress-pattern path still consumes numeric values and recent stress history. Reducing the final score is not equivalent to removing untrusted contributions from every explanation. '
'The 200 missing-signal cases all receive the Normal category despite zero usable physiological channels. Their ready profile reflects the preceding good history, not valid current measurements. '
'The appropriate operational interpretation is unavailable current evidence; a categorical Normal display can conceal that distinction.')
para('Every fall-inactive input reaches exactly 82 and Critical Review because the explicit minimum dominates the weighted score. This consistency confirms execution of the override. '
'It does not establish that a true fall has been detected, since the experiment supplies the fall flag directly. No disease-classification accuracy, sensitivity, specificity, or receiver-operating-characteristic area is calculated.')
h2('B. Impact capture and critical flags')
table('TABLE III. CAPTURE / CRITICAL COUNTS OUT OF 200',['Trace family','Orig. 1 Hz','Orig. 50 Hz','Guard 50 Hz'],[
    [label]+[f"{next(r for r in motion if r['group']==g and r['policy']==p)['captured']} / {next(r for r in motion if r['group']==g and r['policy']==p)['critical']}" for p in policies]
    for g,label in zip(groups,['Quiet','Walking','Impact immobile','Impact recovery','Impact HR drop','Unworn impact'])], [1.13,.79,.79,.79])
para('At 1 Hz, the original function captures 14 of 200 immobile-impact traces (7.0%) and generates zero critical flags. At 50 Hz it captures 199 (99.5%) but still generates zero critical flags. '
'The reason is the cancellation branch: as the impulse decays toward 1 g, the inter-sample difference exceeds 0.10 g while the magnitude is below 2.6 g. The timer is therefore canceled by the settling transient. '
'Increased sampling alone does not repair this state-transition rule.')
figure('motion_results','Fig. 4. Critical-flag rate for each synthetic family and policy. The unworn-impact result is an adverse false-escalation case, not a successful detection.')
para('The guarded candidate escalates 199 of 200 immobile-impact traces (99.5%; Wilson interval 97.22-99.91%) and none of the recovery traces (0%; upper Wilson bound 1.88%). '
'However, it also escalates 198 of 200 unworn-device impacts (99.0%; interval 96.43-99.73%). Both the immobile and unworn families share the same observable motion model. '
'The candidate therefore demonstrates timer preservation without demonstrating reliable discrimination of human emergencies. A single aggregate accuracy would obscure this failure.')
para('For the heart-rate-step family, the original 1 Hz policy produces nine immediate critical flags. The original 50 Hz policy produces none, although it captures 199 impacts. '
'At the higher rate, the heart-rate step can be observed before magnitude crosses the impact threshold; by the time the countdown starts, the adjacent-sample heart-rate difference is zero. '
'The guarded candidate eventually escalates 199 cases after its countdown, rather than restoring immediate escalation. This result exposes an event-association problem: faster acquisition can reveal timing assumptions '
'that are hidden by coarse sampling. A recent-event buffer would be a separate design change requiring its own evaluation.')
h2('C. Controller and backend observations')
elapsed=next(r['observed_elapsed_s'] for r in control if r['case']=='fall_then_silence_32s')
para(f'The production mobile controller triggers after {elapsed:.3f} seconds in the single fall-then-silence run. Movement at five seconds cancels the countdown, manual reset returns it to Idle, '
'and stationary input without a fall remains Idle. A Critical fall activity triggers immediately and is labeled with the sudden-heart-rate-drop reason even though that case supplies no preceding heart-rate decline. '
'Thus, both stale-input handling and trigger-reason provenance require attention. Continued packet silence is currently treated as no-response continuity, rather than distinguished from sensor disconnection.')
para('The backend returns success when either provider stub accepts, and false when both reject. Three synthetic medical facilities appear in each response. This confirms the programmed Boolean aggregation and '
'facility-list path. The failed-delivery webhook is acknowledged with an ok/received response, but the handler maintains no delivery record or alert-state transition. '
'Provider acceptance, handset delivery, and human acknowledgment are therefore separate claims; only the first is represented by these tested return values.')

h1('VI. DISCUSSION AND DESIGN IMPLICATIONS')
para('The experiments show why a safety pipeline should be evaluated at its interfaces. A realistic-looking impact flag, a component score labeled as probability, and a successful backend response each describe '
'a narrower event than the user may infer. In the present prototype, a firmware transient cancels a countdown, missing data maps to Normal, and delivery-like fields describe provider acceptance. '
'These observations are reproducible engineering defects or ambiguities, not statistical proof of clinical failure rates.')
para('The first design implication is to decouple motion acquisition and local event persistence from the connected BLE notification loop. The firmware should retain a bounded event record and expose timestamped '
'states across reconnection. The guarded experiment supports investigating recovery evidence over a time window, but its unworn-impact false escalations show that contact and wearer-state evidence must also be studied. '
'A brief acceleration change alone cannot establish recovery, and prolonged absence of movement alone cannot establish unconsciousness.')
para('The second implication is to make unavailable evidence a first-class output. Quality should gate each contributing channel before scoring, and the final display should distinguish an ordinary supported result '
'from insufficient data. A valid historical baseline must not imply that a missing current reading is valid. Baseline adaptation also requires separate treatment of stationary rest, sleep, exercise, and recovery. '
'The sample-count rules in the current implementation should be replaced or supplemented by duration, coverage, and within-context stability checks evaluated on longitudinal recordings.')
para('The third implication is to model emergency communication as explicit states: created, locally queued, backend accepted, provider accepted, delivered, acknowledged, and resolved or failed. '
'An event identifier and timestamped status history would allow retries and guardian acknowledgment to refer to the same incident. Webhook authentication and persistent status handling are prerequisites for trusting '
'external status updates. These recommendations arise from the inspected contract; they have not been implemented or benchmarked in this study.')
para('Finally, the ECG and stress outputs require a defined measurement and validation pathway. A single periodic ADC-derived ECG scalar cannot support waveform morphology analysis. '
'The GSR scale is not calibrated conductance, and supplied HRV is not equivalent to a quality-controlled interval series. A future measurement study must specify sensor placement, reference equipment, sampling rates, '
'filtering, artifact handling, and inclusion criteria before drawing conclusions about physiology. NeuroTwin is currently most defensible as an interpretable software profile whose assumptions remain visible.')

h1('VII. LIMITATIONS AND VALIDATION PLAN')
para('Synthetic distributions were chosen by the experiment designer and are intentionally simple. They contain no real fall biomechanics, skin-tone variation, age distribution, pathology, or device-position variation. '
'The same seed structure also supports paired comparisons rather than independent clinical cohorts. The guarded policy was introduced during exploratory debugging; there is no held-out policy-selection set. '
'Its high immobile-impact escalation rate is conditioned on the generator and must not be advertised as fall-detection accuracy.')
para('The native test bypasses sensors and radio communication. The Flutter tests bypass BLE parsing and background operating-system restrictions. Backend tests bypass HTTP transport and real provider behavior. '
'The experiments therefore do not constitute an end-to-end validation of a functioning wristband. No battery, memory, energy, notification-delivery probability, GPS accuracy, or ambulance response time is measured. '
'The one controller duration measures a desktop software callback under one execution environment and should not be generalized to a locked Android phone.')
para('External validation should proceed in stages. First, deterministic regression tests should cover the demonstrated counterexamples, stale packets, timestamp order, and missing channels. '
'Second, instrumented bench measurements should establish acquisition timing, signal integrity, packet handling, contact changes, and power use with the selected sensors. '
'Third, ethically reviewed human studies should compare measurements with appropriate references and evaluate ordinary activity before controlled fall-related testing. Any prospective field trial must define '
'consent, monitoring responsibility, stopping rules, and independent event adjudication.')
para('For future dataset modeling, participant-level separation is needed to prevent the same person or overlapping windows appearing in training and evaluation sets. Normalization and feature selection must be '
'fitted using training partitions only. MIT-BIH can support an ECG-specific benchmark after waveform acquisition is implemented; wearable stress resources can support context-dependent research after channel '
'harmonization. These datasets describe different tasks and populations and should not be concatenated into a single unspecified disease label. Calibration and external validation are needed before using '
'probability terminology or offering condition-specific conclusions.')

h1('VIII. CONCLUSION')
para('This study provides a reproducible software evaluation of NeuroGuardian X using actual prototype decision code, controlled synthetic inputs, and isolated notification contracts. '
'The results identify sampling-related impact misses, timer cancellation by post-impact settling, event-timing sensitivity in the heart-rate-drop branch, exclusion of stationary baseline samples, '
'and misleading Normal labels under missing input. An exploratory recovery guard preserves the timeout for most constructed immobile impacts but also escalates nearly all unworn-device impacts. '
'The work therefore contributes testable failure cases and a bounded comparison rather than a clinically validated detector. Completing the measurement pipeline and evaluating these cases on real hardware '
'are necessary next steps before claims about user safety or medical performance.')

h1('DATA AND CODE AVAILABILITY')
para('The accompanying reproducibility package contains the synthetic generators, native extraction harness, Flutter experiment file, offline backend test, trial-level CSV/JSON results, summary tables, '
'editable manuscript sources, and figure PDFs. A SHA-256 manifest identifies the evaluated source files. No participant data, access tokens, guardian contact information, or production environment files are '
'included. The package accompanies this manuscript locally; no public archival DOI is claimed. Authors should deposit an approved, licensed version in a persistent repository before submission.')
h1('ETHICS STATEMENT')
para('This work used generated values and isolated software execution only. No human participants were recruited, no identifiable patient recordings were processed, and no clinical intervention was performed; '
'therefore no human-subject approval was sought for these simulations.')
h1('ACKNOWLEDGMENT')
para('OpenAI Codex using GPT-6 assisted in drafting Sections I-VIII and the declarations, creating the experiment and document-generation code, and generating the plots from executable results. '
'The numerical findings originate from the accompanying experiment outputs rather than invented patient observations. The named authors remain responsible for verifying the code, sources, interpretation, and '
'final manuscript before submission.')

references=[
('A. Sucerquia, J. D. Lopez, and J. F. Vargas-Bonilla, "SisFall: A fall and movement dataset," Sensors, vol. 17, no. 1, Art. no. 198, 2017, doi: 10.3390/s17010198.', 'https://doi.org/10.3390/s17010198'),
('A. Hongn, F. Bosch, L. E. Prado, J. M. Ferrandez, and M. P. Bonomini, "Wearable physiological signals under acute stress and exercise conditions," Scientific Data, vol. 12, Art. no. 520, 2025, doi: 10.1038/s41597-025-04845-9.', 'https://doi.org/10.1038/s41597-025-04845-9'),
('G. B. Moody and R. G. Mark, "The impact of the MIT-BIH Arrhythmia Database," IEEE Engineering in Medicine and Biology Magazine, vol. 20, no. 3, pp. 45-50, 2001, doi: 10.1109/51.932724.', 'https://doi.org/10.1109/51.932724'),
('G. Moody and R. Mark, "MIT-BIH Arrhythmia Database," PhysioNet, version 1.0.0, 2005. Accessed: Sep. 17, 2026. [Online]. Available: https://physionet.org/content/mitdb/1.0.0/', 'https://physionet.org/content/mitdb/1.0.0/'),
('M. R. Amin, D. Wickramasuriya, and R. T. Faghih, "A wearable exam stress dataset for predicting cognitive performance in real-world settings," PhysioNet, version 1.0.0, 2022, doi: 10.13026/kvkb-aj90.', 'https://doi.org/10.13026/kvkb-aj90'),
('PhysioNet, "Sudden Cardiac Death Holter Database," version 1.0.0. Accessed: Sep. 17, 2026. [Online]. Available: https://physionet.org/content/sddb/1.0.0/', 'https://physionet.org/content/sddb/1.0.0/'),
('A. Hongn, F. Bosch, L. Prado, and P. Bonomini, "Wearable device dataset from induced stress and structured exercise sessions," PhysioNet, version 1.0.1, 2025, doi: 10.13026/he0v-tf17.', 'https://doi.org/10.13026/he0v-tf17'),
('Google, "Nearby Search (New)," Google Maps Platform documentation. Accessed: Sep. 17, 2026. [Online]. Available: https://developers.google.com/maps/documentation/places/web-service/nearby-search', 'https://developers.google.com/maps/documentation/places/web-service/nearby-search')]
h1('REFERENCES')
for i,(r,url) in enumerate(references,1):blocks.append({'type':'ref','text':f'[{i}] {r}','url':url})
(ROOT/'manuscript_content.json').write_text(json.dumps({'title':TITLE,'authors':AUTHORS,'abstract':ABSTRACT,'keywords':KEYWORDS,'blocks':blocks},indent=2))

doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75)
sec.left_margin=Inches(.625);sec.right_margin=Inches(.625)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
styles=doc.styles
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    styles[sty].font.name='Times New Roman';styles[sty].font.color.rgb=RGBColor(0,0,0)
    rpr=styles[sty].element.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for key in list(fonts.attrib):
            if 'Theme' in key:del fonts.attrib[key]
    ppr=styles[sty].element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
normal=styles['Normal'];normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=Pt(11.5);normal.paragraph_format.space_after=Pt(2)
normal.paragraph_format.first_line_indent=Inches(.12)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for sty in ['Heading 1','Heading 2']:
    styles[sty].font.size=Pt(10);styles[sty].paragraph_format.first_line_indent=Inches(0)
    styles[sty].paragraph_format.space_before=Pt(8);styles[sty].paragraph_format.space_after=Pt(4)
    styles[sty].paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Heading 1'].font.bold=False
styles['Heading 2'].font.bold=False;styles['Heading 2'].font.italic=True
styles['Caption'].font.size=Pt(8);styles['Caption'].font.italic=False
styles['Caption'].font.bold=False
styles['Caption'].paragraph_format.line_spacing=Pt(9)
styles['Caption'].paragraph_format.first_line_indent=Inches(0)
styles['Caption'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(22);styles['Title'].font.bold=False
styles['Title'].paragraph_format.line_spacing=Pt(25)
styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Title'].paragraph_format.first_line_indent=Inches(0)
p=doc.add_paragraph(TITLE,'Title');p.paragraph_format.space_after=Pt(9)
for val in [AUTHORS,AFFIL,EMAIL]:
    p=doc.add_paragraph(val);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(2)
    for run in p.runs:run.font.size=Pt(10)
sec=doc.add_section(WD_SECTION_START.CONTINUOUS)
cols=sec._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360')
p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0)
p.add_run('Abstract - ').bold=True;p.add_run(ABSTRACT)
for r in p.runs:r.font.size=Pt(9);r.bold=True
p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0)
p.add_run('Index Terms - ').bold=True;p.add_run(KEYWORDS)
for r in p.runs:r.font.size=Pt(9)

def mr(value):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r
def math_container(tag,children):
    node=OxmlElement('m:'+tag)
    for x in children:node.append(x)
    return node
def frac(a,b):return math_container('f',[math_container('num',a),math_container('den',b)])
def sub(a,b):return math_container('sSub',[math_container('e',[mr(a)]),math_container('sub',[mr(b)])])
def sup(children,power):return math_container('sSup',[math_container('e',children),math_container('sup',[mr(power)])])
def radical(children):
    prop=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');prop.append(hide)
    return math_container('rad',[prop,math_container('deg',[]),math_container('e',children)])
def mathline(nodes):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
    p._p.append(math_container('oMath',nodes))
    return p

for b in blocks:
    typ=b['type']
    if typ in ['p','h1','h2','ref']:
        p=doc.add_paragraph(b['text'],{'h1':'Heading 1','h2':'Heading 2'}.get(typ,'Normal'))
        if typ=='ref':
            p.paragraph_format.first_line_indent=Inches(-.18);p.paragraph_format.left_indent=Inches(.18)
            p.paragraph_format.line_spacing=Pt(9);p.paragraph_format.space_after=Pt(4)
            for r in p.runs:r.font.size=Pt(8)
    elif typ=='eq':
        mu=chr(956);sigma=chr(931);pi=chr(960);phi=chr(966)
        if '(1)' in b['text']:
            mathline([mr(mu+' = '),frac([mr('1')],[mr('n')]),mr(sigma+' '),sub('x','i')]).paragraph_format.keep_with_next=True
            mathline([mr('s = '),radical([frac([mr(sigma+' '),sup([mr('('),sub('x','i'),mr(' - '+mu+')')],'2')],[mr('n - 1')])]),mr('   (1)')])
        elif '(2)' in b['text']:
            mathline([mr('z = '),frac([sub('x','current'),mr(' - '+mu)],[mr('max(s, '),sub('s','min'),mr(')')]),mr('   (2)')])
        elif '(3)' in b['text']:
            mathline([sub('R','0'),mr(' = 0.18S + 0.22C + 0.18E')]).paragraph_format.keep_with_next=True
            mathline([mr('+ 0.26F + 0.10O + 0.06T   (3)')])
        else:
            mathline([mr('a(t) = 1 + 0.004 sin(2'+pi+' 0.7t + '+phi+')')]).paragraph_format.keep_with_next=True
            mathline([mr('+ (A - 1) max(0, 1 - '),frac([mr('2|t - '),sub('t','0'),mr('|')],[mr('w')]),mr(')   (4)')])
    elif typ=='fig':
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Inches(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(FIG/(b['name']+'.png')),width=Inches(3.48))
        doc.add_paragraph(b['caption'],'Caption')
    elif typ=='table':
        p=doc.add_paragraph(b['caption'],'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        t=doc.add_table(rows=1,cols=len(b['head']));t.autofit=False
        for cell,width,title in zip(t.rows[0].cells,b['widths'],b['head']):cell.width=Inches(width);cell.text=title
        for row in b['rows']:
            cells=t.add_row().cells
            for cell,width,val in zip(cells,b['widths'],row):cell.width=Inches(width);cell.text=str(val)
        for col,width in zip(t.columns,b['widths']):col.width=Inches(width)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            tag=OxmlElement('w:'+side);tag.set(qn('w:val'),'single');tag.set(qn('w:sz'),'4');tag.set(qn('w:color'),'D9D9D9');borders.append(tag)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),'55');e.set(qn('w:type'),'dxa');marg.append(e)
                cp.append(marg)
                if i==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EDEFF1');cp.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(9)
                    p.paragraph_format.keep_with_next=(i<len(t.rows)-1)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:r.font.size=Pt(8);r.bold=(i==0)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)

for section in doc.sections[:1]:
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    foot.paragraph_format.first_line_indent=Inches(0)

# Remove inherited decorative paragraph borders and theme font overrides.
for style in styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for fonts in style.element.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if 'Theme' in key:del fonts.attrib[key]
doc.core_properties.title=TITLE;doc.core_properties.author='';doc.core_properties.subject='Synthetic software evaluation of NeuroGuardian X'
doc.core_properties.comments='Editable IEEE-style journal manuscript. Author details require completion.'
doc.save(ROOT/'NeuroGuardian_X_IEEE_Manuscript.docx')

def texesc(s):
    for a,b in [('\\',r'\textbackslash{}'),('&',r'\&'),('%',r'\%'),('$',r'\$'),('#',r'\#'),('_',r'\_')]:s=s.replace(a,b)
    return s
tex=[r'\documentclass[journal]{IEEEtran}',r'\usepackage{graphicx,amsmath,array,url}',
     r'\graphicspath{{figures/}}',r'\begin{document}',r'\title{'+texesc(TITLE)+'}',
     r'\author{'+texesc(AUTHORS)+r'\thanks{'+texesc(AFFIL+'; '+EMAIL)+r'}}',r'\maketitle',
     r'\begin{abstract}'+texesc(ABSTRACT)+r'\end{abstract}',r'\begin{IEEEkeywords}'+texesc(KEYWORDS)+r'\end{IEEEkeywords}']
inrefs=False
for b in blocks:
    typ=b['type']
    if typ=='h1':
        title=b['text']
        if title=='REFERENCES':tex.append(r'\begin{thebibliography}{8}');inrefs=True
        elif title[0] in 'IVX' and '. ' in title:tex.append(r'\section{'+texesc(title.split('. ',1)[1].title())+'}')
        else:tex.append(r'\section*{'+texesc(title.title())+'}')
    elif typ=='h2':tex.append(r'\subsection{'+texesc(b['text'].split('. ',1)[1])+'}')
    elif typ=='p':tex.append(texesc(b['text'])+'\n')
    elif typ=='eq':tex.append(r'\begin{equation}'+b['latex']+r'\end{equation}')
    elif typ=='fig':
        cap=b['caption'].split('. ',2)[2]
        tex.append(r'\begin{figure}[!t]\centering\includegraphics[width=\columnwidth]{'+b['name']+r'.pdf}\caption{'+texesc(cap)+r'}\end{figure}')
    elif typ=='table':
        spec=''.join('p{'+f'{w/3.5*.91:.3f}'+r'\columnwidth}' for w in b['widths'])
        cap=b['caption'].split('. ',1)[1]
        tex.extend([r'\begin{table}[!t]\caption{'+texesc(cap)+r'}\centering\scriptsize\setlength{\tabcolsep}{2pt}',r'\begin{tabular}{'+spec+r'}\hline',
                    ' & '.join(texesc(x) for x in b['head'])+r'\\\hline'])
        tex += [' & '.join(texesc(str(x)) for x in row)+r'\\' for row in b['rows']]
        tex.append(r'\hline\end{tabular}\end{table}')
    elif typ=='ref':
        n=b['text'].split(']')[0][1:]
        raw=b['text'].split('] ',1)[1]
        chunks=re.split(r'(https?://\S+)',raw)
        rendered=''.join(r'\url{'+x+'}' if x.startswith('http') else texesc(x) for x in chunks)
        tex.append(r'\bibitem{ref'+n+'}'+rendered)
if inrefs:tex.append(r'\end{thebibliography}')
tex.append(r'\end{document}')
(ROOT/'NeuroGuardian_X_IEEE_Manuscript.tex').write_text('\n'.join(tex),encoding='utf-8')

sourcefiles=['pubspec.yaml','pubspec.lock','firmware/esp32_neuroguardian_ble/esp32_neuroguardian_ble.ino',
 'lib/features/metrics/metrics.dart','lib/features/metrics/metrics_providers.dart','lib/features/metrics/stress_analysis.dart',
 'lib/features/emergency/auto_sos_provider.dart','backend/emergency_server.py']
sourcefiles += [str(p.relative_to(APP)).replace('\\','/') for p in (APP/'lib/features/neurotwin').glob('*.dart')]
manifest=[]
for f in sourcefiles:
    p=APP/f;dst=ROOT/'source_snapshot'/f;dst.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dst)
    manifest.append({'path':f,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(ROOT/'source_manifest.json').write_text(json.dumps({'snapshot_date':'2026-09-17','app_version':'0.4.0+4',
 'git_base_commit':'7215f020a354033096998de0c285f8ba2eacaab5','working_tree_modified':True,
 'files':manifest,'python_artifact_runtime':sys.version.split()[0],
 'training_records':0,'patient_records':0,'engine_evaluations':1600,'motion_traces':1200,
 'motion_policy_evaluations':3600,'controller_cases':5,'backend_cases':5},indent=2))

notes='''# NeuroGuardian X manuscript and reproducibility package

## Files
- NeuroGuardian_X_IEEE_Manuscript.docx: editable two-column Word manuscript.
- NeuroGuardian_X_IEEE_Manuscript.pdf: visual copy exported from the Word file.
- NeuroGuardian_X_IEEE_Manuscript.tex: editable IEEEtran journal source; compile with pdfLaTeX in an installation containing IEEEtran, graphicx, amsmath, array and url (for example Overleaf).
- figures/: vector PDF figures and print-resolution PNGs.
- results/: actual trial-level outputs and summaries.
- simulation/: runnable experiments, including an extracted C++ function.
- source_snapshot/ and source_manifest.json: inspected implementation files and hashes. This is an audit snapshot, not a complete Flutter application.

## Scope
All physiological and motion inputs are synthetic. No PhysioNet dataset or human-subject recording was used; no model was trained. No alert was transmitted. The paper is a complete simulation-study manuscript requiring author review, not evidence of clinical validation or acceptance by a journal.

The experimental guarded recovery policy exists only in the simulation. It is not installed on the watch or phone. Its unworn-device false alerts and missed immediate heart-rate escalation are reported, not hidden.

## Author and submission tasks
1. Fill the author names, affiliations, and corresponding email in the editable title block. No identities or IEEE memberships were invented.
2. Choose the target journal and transfer to its current required template if it differs from a general IEEE journal layout.
3. Review every result, interpretation, reference, and AI disclosure. Resolve funding, conflicts, contributions, and code licensing as appropriate.
4. Extend the evidence with hardware and external validation if required by the journal and the intended claims. A synthetic-only prototype failure study has limited generalizability and may not satisfy a journal's novelty threshold.
5. Deposit an approved reproducibility archive and add its persistent link. Nothing in this package has been submitted or publicly published.

IEEE template guidance: https://journals.ieeeauthorcenter.ieee.org/create-your-ieee-journal-article/authoring-tools-and-templates/
IEEE AI and human-subject disclosure policy: https://journals.ieeeauthorcenter.ieee.org/become-an-ieee-journal-author/publishing-ethics/guidelines-and-policies/submission-and-peer-review-policies/
IEEEtran class: https://ctan.org/pkg/IEEEtran

## Reproduction on the original project machine
The package is stored in <workspace>/output/NeuroGuardianX_IEEE_Research_2026-09-17 and the app in <workspace>/neuroguardian_app.

1. Run simulation/run_motion_experiment.py with Python. On this Windows machine its build step uses the installed MSVC vcvars64.bat. On another machine, replace that build step with a C++ compiler invocation for generated native_motion.cpp. There are no third-party Python requirements for this simulation.
2. From the Flutter app root, set NGX_RESEARCH_RESULTS to the package results directory and run flutter test --no-pub <absolute-package-path>/simulation/neurotwin_experiment_test.dart --reporter expanded. Flutter dependencies must already be installed. The timer experiment takes 32 seconds and uses real elapsed time. Repeated controller latency will vary.
3. Run simulation/backend_contract_experiment.py with the backend Python environment containing FastAPI and HTTPX. It stubs providers, disables dotenv, and blocks network clients. Do not remove those safeguards to reproduce the offline results.
4. Run build_paper.py with Python packages python-docx and reportlab and Poppler available. The artifact script currently specifies this machine's Poppler path. It regenerates tables, figures, DOCX and LaTeX from the measured output files.

## Counts and interpretation
Engine: 8 scenarios x 5 seeds x 40 repetitions = 1600 evaluations. Motion: 6 families x 5 seeds x 40 repetitions = 1200 traces, each replayed under 3 policies = 3600 evaluations. Controller: 5 local cases. Backend: 4 provider-outcome combinations plus 1 failed-delivery webhook.

Percentages describe generated traces, not patients. Labels such as Normal, fall probability, and ECG probability are implementation labels, not medically verified outcomes. Current code returns a Normal category with unavailable vitals; this is a finding to correct, not a feature to advertise.

No secret .env, guardian number, production access token, or real location is distributed. Synthetic backend coordinates are arbitrary test values.
'''
(ROOT/'README.md').write_text(notes,encoding='utf-8')
wordcount=len((ABSTRACT+' '+ ' '.join(b.get('text','') for b in blocks)).split())
print(json.dumps({'docx':str(ROOT/'NeuroGuardian_X_IEEE_Manuscript.docx'),'words_including_references':wordcount,'figures':4,'tables':3}))
