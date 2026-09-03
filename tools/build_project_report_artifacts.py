from __future__ import annotations

import sys
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
DOCX_OUT = DOCS_DIR / "NeuroGuardian_X_Project_Report.docx"
PDF_OUT = DOCS_DIR / "NeuroGuardian_X_Project_Report.pdf"

SKILL_ROOT = Path(
    r"C:\Users\hp\.codex\plugins\cache\openai-primary-runtime\documents\26.826.12353\skills\documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry, exact_column_widths  # noqa: E402


PREPARED_DATE = "September 2, 2026"
REPORT_VERSION = "0.3.0+3"


NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
TEAL = RGBColor(14, 124, 123)
GRAY = RGBColor(82, 91, 108)
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
MINT = "E6F5F2"
WARM = "FFF8E8"
RISK = "FCE8E6"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)


def dxa_from_inches(value: float) -> int:
    return int(round(value * 1440))


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 6, line: float = 1.10) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9DEE8", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=GRAY)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r_begin = OxmlElement("w:r")
    r_begin.append(fld_begin)
    r_instr = OxmlElement("w:r")
    r_instr.append(instr)
    r_end = OxmlElement("w:r")
    r_end.append(fld_end)

    paragraph._p.append(r_begin)
    paragraph._p.append(r_instr)
    paragraph._p.append(r_end)


def add_bottom_border(paragraph, color: str = "0E7C7B", size: str = "10") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(8)
    bullet.paragraph_format.line_spacing = 1.167

    number = styles["List Number"]
    number.font.name = "Calibri"
    number.font.size = Pt(11)
    number.paragraph_format.left_indent = Inches(0.5)
    number.paragraph_format.first_line_indent = Inches(-0.25)
    number.paragraph_format.space_after = Pt(8)
    number.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "NeuroGuardian X Project Report"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.runs[0], size=9, color=GRAY, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_para(
    doc: Document,
    text: str,
    *,
    style: str | None = None,
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    align: WD_ALIGN_PARAGRAPH | None = None,
    before: float = 0,
    after: float = 6,
    line: float = 1.10,
) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    set_paragraph_spacing(p, before, after, line)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def add_list(doc: Document, items: Iterable[str], *, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        set_paragraph_spacing(p, 0, 8, 1.167)
        run = p.add_run(item)
        set_run_font(run, size=11, color=BLACK)


def add_callout(doc: Document, title: str, body: str, fill: str = MINT) -> None:
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, color="D7DBE2")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, 0, 2, 1.10)
    r = p.add_run(title.upper())
    set_run_font(r, size=9, color=TEAL, bold=True)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, 0, 0, 1.10)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=NAVY)
    add_para(doc, "", after=4)


def add_data_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_in: Sequence[float],
    *,
    header_fill: str = BLUE_GRAY,
    font_size: float = 9.2,
) -> None:
    widths = exact_column_widths([dxa_from_inches(w) for w in widths_in], 9360)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    mark_header_row(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade_cell(cell, header_fill)
        set_cell_border(cell, color="CDD6E3")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, 0, 0, 1.10)
        r = p.add_run(text)
        set_run_font(r, size=9.2, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_border(cell, color="DEE3EA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, 0, 0, 1.10)
            r = p.add_run(text)
            set_run_font(r, size=font_size, color=BLACK)
    apply_table_geometry(table, widths, indent_dxa=120)
    add_para(doc, "", after=4)


def add_code_block(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    apply_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    shade_cell(cell, "0B2545")
    set_cell_border(cell, color="0B2545")
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, 0, 0, 1.0)
    r = p.add_run(text)
    set_run_font(r, name="Courier New", size=8.2, color=WHITE)
    add_para(doc, "", after=4)


def build_docx() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc)

    # First page uses the memo_masthead pattern with a concise technical report stack.
    add_para(doc, "PROJECT REPORT / WEARABLE SAFETY SYSTEM", size=9.5, color=TEAL, bold=True, after=2)
    add_para(doc, "NeuroGuardian X", size=28, color=NAVY, bold=True, after=2)
    add_para(
        doc,
        "Project overview, input data, output data, architecture, current status, and release-readiness report.",
        size=13.5,
        color=GRAY,
        after=14,
    )
    for label, value in [
        ("Prepared date", PREPARED_DATE),
        ("App version", REPORT_VERSION),
        ("Current stage", "Hardware-connected prototype moving toward real-world pilot testing"),
        ("Primary safety goal", "Detect falls and no-movement emergencies, then alert guardians with GPS and nearby hospital context"),
    ]:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, 0, 2, 1.10)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.8, color=NAVY, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.8, color=BLACK)
    divider = doc.add_paragraph()
    add_bottom_border(divider, color="0E7C7B", size="12")
    set_paragraph_spacing(divider, 8, 12, 1.0)

    add_callout(
        doc,
        "Executive position",
        "NeuroGuardian X is an emergency-support and risk-detection system, not a medical diagnosis device. The safest path to society-scale use is a phased pilot: real sensors, calibrated thresholds, verified guardian notifications, emergency-service partnerships, and formal privacy/medical review.",
        MINT,
    )

    add_heading = doc.add_heading
    add_heading("1. Project Overview", level=1)
    add_para(
        doc,
        "NeuroGuardian X combines an ESP32-S3 wrist-worn device, a Flutter Android app, and a small emergency backend. The watch streams movement and health signals to the phone over Bluetooth Low Energy. The app detects risk patterns, starts a 30-second fall countdown, lets the user cancel with an I AM OKAY control, and sends emergency alerts when the situation remains critical.",
    )
    add_list(
        doc,
        [
            "Target users: older adults, workers, patients under observation, and people who may need rapid guardian awareness after a fall.",
            "Main value: faster response during fall, no-movement, sudden heart-rate drop, low oxygen, stress, temperature, or manual SOS conditions.",
            "Important boundary: the system supports emergency awareness; it should not claim heart attack confirmation, diagnosis, or replacement for a doctor.",
        ],
    )

    add_heading("2. What Has Been Built", level=1)
    add_data_table(
        doc,
        ["Area", "Current implementation", "Real-world role"],
        [
            ["Flutter app", "Android APK installed for phone testing; includes dashboard, device, SOS, care, history, and settings flows.", "Main user interface and emergency control layer."],
            ["BLE connection", "App scans for NeuroGuardianX and listens for JSON packets from the ESP32-S3 wearable.", "Live connection between wrist hardware and phone."],
            ["Fall protocol", "Fall signal starts a 30-second timer; movement cancels it; critical fall can trigger immediate SOS.", "Core safety automation."],
            ["Wearable commands", "App can send BUZZER_ON and BUZZER_OFF to hardware during alerts.", "Bidirectional alert feedback."],
            ["Backend", "FastAPI backend keeps Google/WhatsApp keys out of app code and prepares alert payloads.", "Secure notification and hospital lookup layer."],
            ["Notifications", "Google Places lookup and ntfy push fallback are working; WhatsApp waits on Meta template approval.", "Guardian and emergency awareness."],
            ["Firmware", "ESP32-S3 firmware has real-sensor mode, IMU support, optional MAX30102/GPS flags, buzzer, vibration, button, and battery hooks.", "Wearable data producer."],
        ],
        [1.25, 3.35, 1.9],
    )

    add_heading("3. System Architecture", level=1)
    add_para(doc, "The system is split into three layers: wearable sensing, phone-side decision UX, and backend alert routing.")
    add_data_table(
        doc,
        ["Layer", "Input", "Processing", "Output"],
        [
            ["ESP32-S3 watch", "IMU, heart sensor, SpO2, GSR, temperature, ECG, GPS, button, battery", "Reads sensors, detects impact/no movement, streams compact JSON over BLE", "Vitals packets, fall events, buzzer/vibration feedback"],
            ["Flutter app", "BLE packets, user settings, guardian number, phone GPS", "Risk detection, fall countdown, movement cancellation, care suggestions, alert history", "Live dashboard, I AM OKAY reset, SOS request, local warnings"],
            ["Emergency backend", "SOS payload, location, guardian target, API keys stored in .env", "Hospital lookup, message formatting, WhatsApp/push routing", "Guardian alert with Google Maps link and top medical facilities"],
        ],
        [1.3, 1.75, 2.35, 1.1],
    )
    add_code_block(
        doc,
        "ESP32-S3 Watch -> BLE JSON -> Flutter App -> 30-sec Safety Logic -> Backend -> Google Places -> WhatsApp / Push / SMS fallback"
    )

    doc.add_page_break()
    add_heading("4. Input Data", level=1)
    add_heading("4.1 Wearable Sensor Inputs", level=2)
    add_data_table(
        doc,
        ["Input", "Suggested source", "Purpose", "Current status"],
        [
            ["Acceleration and gyro", "MPU6050 now; BMI270/ICM-42688 recommended for production", "Fall detection, movement recovery, activity context", "Firmware hooks present"],
            ["Heart rate", "MAX30102/MAX86141 optical PPG", "Sudden HR drop, cardiac load, stress support", "Optional firmware flag"],
            ["SpO2", "MAX30102/MAX86141 optical PPG", "Low oxygen risk detection", "Optional sensor calibration needed"],
            ["HRV", "Derived from beat-to-beat interval", "Stress and cardiac trend support", "Model-ready input"],
            ["GSR/EDA", "Skin conductance sensor", "Stress response detection", "Analog input present"],
            ["Body temperature", "MAX30205/MAX30208 or skin temp sensor", "Fever/heat strain context", "App temperature meter present"],
            ["ECG amplitude", "AD8232 or compact ECG front-end", "Future ECG model input", "Analog input present; medical validation needed"],
            ["GPS", "Watch GPS module or phone GPS fallback", "SOS location link and nearby facility lookup", "Backend and app payload support present"],
            ["Battery", "Voltage divider/fuel gauge", "Low battery risk warning", "Firmware input present"],
            ["Manual SOS", "Watch button", "User-triggered emergency/test path", "Firmware input present"],
        ],
        [1.25, 1.75, 2.3, 1.2],
        font_size=8.6,
    )

    add_heading("4.2 User and Backend Inputs", level=2)
    add_list(
        doc,
        [
            "Guardian WhatsApp number is entered by each wearable owner; it is not hard-coded by the developer.",
            "Backend URL, device token, and push target connect the app to the alert service.",
            "Google Maps and WhatsApp provider secrets stay in the backend .env file, not inside Flutter app code.",
            "Doctor-suggested medicines are stored as prescriptions entered by a user/doctor; the app should not automatically prescribe medicine.",
        ],
    )

    doc.add_page_break()
    add_heading("4.3 Dataset Inputs for Future AI Model", level=2)
    add_data_table(
        doc,
        ["Dataset", "Signal focus", "Planned use"],
        [
            ["MIT-BIH Arrhythmia Database", "ECG rhythm examples", "Train and validate arrhythmia pattern research models."],
            ["Sudden Cardiac Death Holter Database", "Long ECG records around serious cardiac events", "Study high-risk ECG features; never use as the only emergency rule."],
            ["Wearable Device Dataset", "BVP, EDA, temperature, motion", "Build wearable stress/activity features similar to watch input."],
            ["Wearable Exam Stress Dataset", "Real-world stress-related wearable signals", "Validate stress trend detection and reduce false stress alerts."],
            ["Prototype watch logs", "Actual device packets from users/tests", "Most important dataset for calibration, fall thresholds, and field performance."],
        ],
        [2.0, 1.65, 2.85],
        font_size=8.8,
    )

    add_heading("5. Output Data", level=1)
    add_data_table(
        doc,
        ["Output", "Destination", "Meaning"],
        [
            ["Live vitals", "Dashboard", "HR, SpO2, HRV, stress, GSR, ECG, body temperature, activity, movement, battery, GPS state."],
            ["AI risk detection label", "Dashboard", "User-friendly condition label such as minor stress, high temperature, low oxygen, or critical fall."],
            ["Fall countdown", "Dashboard and SOS tab", "30-second timer after fall; movement cancels the timer; no movement triggers SOS."],
            ["I AM OKAY reset", "Dashboard and SOS tab", "Stops countdown, turns off wearable buzzer, and records cancellation."],
            ["Care suggestions", "Care tab", "Exercise/home-remedy style guidance for minor criteria; serious criteria shows consult doctor."],
            ["Prescription list", "Care tab", "Doctor-suggested medicines saved by the owner/doctor, not generated as diagnosis."],
            ["Guardian alert", "WhatsApp/push/SMS fallback", "Reason, latest vitals, exact Google Maps link, and top nearby medical facilities."],
            ["Wearable feedback", "Watch buzzer/vibration/LED", "Immediate physical feedback during emergency countdown and connection events."],
        ],
        [1.55, 1.45, 3.5],
        font_size=8.8,
    )

    add_heading("6. Emergency Logic", level=1)
    add_heading("6.1 Fall and No-Movement Workflow", level=2)
    add_list(
        doc,
        [
            "ESP32-S3 detects fall impact or sends a fall_detected event.",
            "The app starts a 30-second emergency countdown with vibration pulses every 2 seconds.",
            "The app sends BUZZER_ON to the watch so the hardware gives local feedback.",
            "If movement returns after the fall, the app stops the timer and sends BUZZER_OFF.",
            "If no movement remains for 30 seconds, the app fetches GPS and creates an SOS alert.",
            "If fall plus sudden heart-rate drop is detected, the app triggers immediate SOS without waiting.",
        ],
        numbered=True,
    )
    add_heading("6.2 SOS Payload", level=2)
    add_code_block(
        doc,
        '{"event_type":"critical_fall","reason":"Fall detected with no movement","location":{"maps_link":"https://maps.google.com/?q=lat,lon"},"medical_facilities":[{"name":"Hospital 1"},{"name":"Hospital 2"},{"name":"Hospital 3"}],"vitals":{"hr":58,"spo2":96,"temperature":36.7,"movement":false}}'
    )

    add_heading("7. App Modules", level=1)
    add_data_table(
        doc,
        ["Module", "Purpose"],
        [
            ["Dashboard", "Eye-catching overview of vitals, temperature meter, risk state, device status, and emergency timer."],
            ["Device", "BLE scan/connect flow, wearable status, packet stream, and reconnect handling."],
            ["SOS", "Manual SOS, emergency countdown state, cancellation, and alert status."],
            ["Care", "Minor condition suggestions, serious-condition doctor warning, and doctor prescription storage."],
            ["History", "Past alert/event timeline for review and testing."],
            ["Settings", "Guardian WhatsApp number, backend URL, device token, ambulance fallback, and notification preferences."],
            ["Safety consent", "Clear disclaimer that this is not a diagnosis device and emergency features need testing."],
        ],
        [1.4, 5.1],
    )

    doc.add_page_break()
    add_heading("8. Hardware and Firmware Architecture", level=1)
    add_para(
        doc,
        "The recommended wrist-watch hardware keeps sensing, feedback, and power control modular so the prototype can be tested sensor by sensor before a custom PCB is produced.",
    )
    add_data_table(
        doc,
        ["Subsystem", "Recommended part", "Reason"],
        [
            ["Controller", "ESP32-S3 module", "BLE support, enough processing power, common dev-board ecosystem."],
            ["Motion", "BMI270/ICM-42688 for production; MPU6050 for low-cost prototype", "Better fall and movement detection with calibrated acceleration/gyro data."],
            ["Optical vitals", "MAX86141/MAX30102", "Heart rate and SpO2 support for wrist prototypes."],
            ["Temperature", "MAX30205/MAX30208 or skin temperature sensor", "Body temperature trend input."],
            ["Stress", "GSR/EDA electrodes", "Stress trend signal when combined with HRV and movement."],
            ["ECG optional", "AD8232 or medical-grade ECG front-end", "Research input only unless clinically validated."],
            ["Location", "Phone GPS fallback plus optional watch GPS", "Best reliability because indoor GPS can fail on small wearables."],
            ["Feedback", "Vibration motor, buzzer, LED", "Local alert and user awareness during countdown."],
            ["Power", "LiPo battery, charger IC, fuel gauge", "Wearable safety depends on reliable power and low-battery warnings."],
        ],
        [1.35, 2.05, 3.1],
        font_size=8.7,
    )

    doc.add_page_break()
    add_heading("9. Release Readiness", level=1)
    add_callout(
        doc,
        "Biggest risk",
        "The product must not miss a real emergency and must not create frequent false emergencies. This is the central safety problem to solve before public release.",
        RISK,
    )
    add_data_table(
        doc,
        ["Risk", "Why it matters", "Control needed"],
        [
            ["False fall alert", "Can worry guardians or waste emergency resources.", "Test with walking, running, sitting, stairs, phone drops, and sleep movement."],
            ["Missed fall", "Can delay help when the user is unconscious.", "Tune IMU thresholds using real falls/safe simulations and multiple body positions."],
            ["Loose sensor contact", "Wrist vitals can be inaccurate.", "Add signal-quality checks and show warnings when readings are unreliable."],
            ["Indoor GPS failure", "Location may be poor inside buildings.", "Use phone GPS, last-known location, Wi-Fi/network location, and accuracy labels."],
            ["Notification failure", "Guardian may not receive message.", "Use WhatsApp, push fallback, SMS composer fallback, retry logging, and delivery status."],
            ["Medical/legal exposure", "Public health products require careful claims.", "Use disclaimers, consent, privacy policy, field trials, and clinical/regulatory advice."],
        ],
        [1.35, 2.15, 3.0],
        font_size=8.6,
    )

    add_heading("10. Next Milestones", level=1)
    add_list(
        doc,
        [
            "Wait for Meta WhatsApp template approval, then test automatic WhatsApp alerts with the owner-entered guardian number.",
            "Wire final IMU, PPG, temperature, GSR, ECG, GPS, battery, button, buzzer, and vibration modules to ESP32-S3.",
            "Enable firmware flags for MAX30102 and GPS only after libraries are installed and wiring is verified.",
            "Run controlled fall/no-movement tests and tune thresholds against false alarms.",
            "Collect prototype logs and train AI models using PhysioNet plus real watch data.",
            "Prepare privacy policy, user consent, emergency-contact onboarding, and pilot-study documentation before public use.",
        ],
    )

    doc.add_page_break()
    add_heading("11. Local Development References", level=1)
    add_data_table(
        doc,
        ["Item", "Location"],
        [
            ["Flutter app", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app"],
            ["Android APK", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\build\app\outputs\flutter-apk\app-debug.apk"],
            ["ESP32 firmware", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\firmware\esp32_neuroguardian_ble\esp32_neuroguardian_ble.ino"],
            ["Backend server", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\backend\emergency_server.py"],
            ["Backend secrets", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\backend\.env (do not publish)"],
            ["AI model plan", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\docs\physionet_model_training_plan.md"],
        ],
        [1.65, 4.85],
        font_size=8.0,
    )

    doc.core_properties.title = "NeuroGuardian X Project Report"
    doc.core_properties.subject = "Project overview, input data, output data, architecture, and release readiness"
    doc.core_properties.author = "NeuroGuardian X"
    doc.core_properties.last_modified_by = "NeuroGuardian X"
    doc.save(DOCX_OUT)


def p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


def bullets(items: Sequence[str], style: ParagraphStyle) -> ListFlowable:
    return ListFlowable(
        [ListItem(p(item, style), leftIndent=14) for item in items],
        bulletType="bullet",
        leftIndent=18,
        bulletFontName="Helvetica",
        bulletFontSize=8,
    )


def pdf_table(data: Sequence[Sequence[str]], widths: Sequence[float], *, header: bool = True) -> Table:
    table = Table([[p(str(cell), PDF_STYLES["TableCell"]) for cell in row] for row in data], colWidths=[w * inch for w in widths], repeatRows=1 if header else 0)
    commands = [
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9DEE8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9DEE8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def draw_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D7DBE2"))
    canvas.setLineWidth(0.5)
    canvas.line(doc.leftMargin, 0.58 * inch, letter[0] - doc.rightMargin, 0.58 * inch)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#525B6C"))
    canvas.drawString(doc.leftMargin, 0.40 * inch, "NeuroGuardian X Project Report")
    canvas.drawRightString(letter[0] - doc.rightMargin, 0.40 * inch, f"Page {doc.page}")
    canvas.restoreState()


BASE_STYLES = getSampleStyleSheet()
PDF_STYLES = {
    "Kicker": ParagraphStyle(
        "Kicker",
        parent=BASE_STYLES["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#0E7C7B"),
        spaceAfter=3,
    ),
    "Title": ParagraphStyle(
        "Title",
        parent=BASE_STYLES["Title"],
        fontName="Helvetica-Bold",
        fontSize=29,
        leading=34,
        textColor=colors.HexColor("#0B2545"),
        alignment=TA_LEFT,
        spaceAfter=6,
    ),
    "Subtitle": ParagraphStyle(
        "Subtitle",
        parent=BASE_STYLES["Normal"],
        fontName="Helvetica",
        fontSize=12.5,
        leading=16,
        textColor=colors.HexColor("#525B6C"),
        spaceAfter=12,
    ),
    "H1": ParagraphStyle(
        "H1",
        parent=BASE_STYLES["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#2E74B5"),
        spaceBefore=14,
        spaceAfter=7,
    ),
    "H2": ParagraphStyle(
        "H2",
        parent=BASE_STYLES["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#1F4D78"),
        spaceBefore=9,
        spaceAfter=5,
    ),
    "Body": ParagraphStyle(
        "Body",
        parent=BASE_STYLES["BodyText"],
        fontName="Helvetica",
        fontSize=9.8,
        leading=13,
        textColor=colors.black,
        spaceAfter=6,
    ),
    "Small": ParagraphStyle(
        "Small",
        parent=BASE_STYLES["BodyText"],
        fontName="Helvetica",
        fontSize=8.3,
        leading=10.3,
        textColor=colors.HexColor("#525B6C"),
        spaceAfter=4,
    ),
    "TableCell": ParagraphStyle(
        "TableCell",
        parent=BASE_STYLES["BodyText"],
        fontName="Helvetica",
        fontSize=7.6,
        leading=9.3,
        textColor=colors.black,
    ),
    "Callout": ParagraphStyle(
        "Callout",
        parent=BASE_STYLES["BodyText"],
        fontName="Helvetica",
        fontSize=9.3,
        leading=12,
        textColor=colors.HexColor("#0B2545"),
        spaceAfter=0,
    ),
}


def callout_flowable(title: str, body: str, fill: str = "#E6F5F2") -> Table:
    data = [[Paragraph(f"<b>{escape(title.upper())}</b><br/>{escape(body)}", PDF_STYLES["Callout"])]]
    table = Table(data, colWidths=[6.5 * inch])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(fill)),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D7DBE2")),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    return table


def build_pdf() -> None:
    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=0.9 * inch,
        bottomMargin=0.85 * inch,
        title="NeuroGuardian X Project Report",
        author="NeuroGuardian X",
    )
    story = []
    story.append(p("PROJECT REPORT / WEARABLE SAFETY SYSTEM", PDF_STYLES["Kicker"]))
    story.append(p("NeuroGuardian X", PDF_STYLES["Title"]))
    story.append(p("Project overview, input data, output data, architecture, current status, and release-readiness report.", PDF_STYLES["Subtitle"]))
    story.append(
        pdf_table(
            [
                ["Prepared date", PREPARED_DATE],
                ["App version", REPORT_VERSION],
                ["Current stage", "Hardware-connected prototype moving toward real-world pilot testing"],
                ["Primary safety goal", "Detect falls and no-movement emergencies, then alert guardians with GPS and nearby hospital context"],
            ],
            [1.65, 4.85],
            header=False,
        )
    )
    story.append(Spacer(1, 10))
    story.append(callout_flowable("Executive position", "NeuroGuardian X is an emergency-support and risk-detection system, not a medical diagnosis device. The safest path to society-scale use is a phased pilot: real sensors, calibrated thresholds, verified guardian notifications, emergency-service partnerships, and formal privacy/medical review."))
    story.append(Spacer(1, 8))

    def h1(text: str) -> None:
        story.append(p(text, PDF_STYLES["H1"]))

    def h2(text: str) -> None:
        story.append(p(text, PDF_STYLES["H2"]))

    h1("1. Project Overview")
    story.append(p("NeuroGuardian X combines an ESP32-S3 wrist-worn device, a Flutter Android app, and a small emergency backend. The watch streams movement and health signals to the phone over Bluetooth Low Energy. The app detects risk patterns, starts a 30-second fall countdown, lets the user cancel with an I AM OKAY control, and sends emergency alerts when the situation remains critical.", PDF_STYLES["Body"]))
    story.append(bullets(["Target users: older adults, workers, patients under observation, and people who may need rapid guardian awareness after a fall.", "Main value: faster response during fall, no-movement, sudden heart-rate drop, low oxygen, stress, temperature, or manual SOS conditions.", "Important boundary: the system supports emergency awareness; it should not claim heart attack confirmation, diagnosis, or replacement for a doctor."], PDF_STYLES["Body"]))

    h1("2. Current Build Status")
    story.append(
        pdf_table(
            [
                ["Area", "Current implementation", "Real-world role"],
                ["Flutter app", "Android APK installed for phone testing; includes dashboard, device, SOS, care, history, and settings flows.", "Main user interface and emergency control layer."],
                ["BLE connection", "App scans for NeuroGuardianX and listens for JSON packets from the ESP32-S3 wearable.", "Live connection between wrist hardware and phone."],
                ["Fall protocol", "Fall signal starts a 30-second timer; movement cancels it; critical fall can trigger immediate SOS.", "Core safety automation."],
                ["Backend", "FastAPI backend keeps Google/WhatsApp keys out of app code and prepares alert payloads.", "Secure notification and hospital lookup layer."],
                ["Firmware", "ESP32-S3 firmware has real-sensor mode, IMU support, optional MAX30102/GPS flags, buzzer, vibration, button, and battery hooks.", "Wearable data producer."],
            ],
            [1.2, 3.35, 1.95],
        )
    )

    h1("3. System Architecture")
    story.append(
        pdf_table(
            [
                ["Layer", "Input", "Processing", "Output"],
                ["ESP32-S3 watch", "IMU, heart sensor, SpO2, GSR, temperature, ECG, GPS, button, battery", "Reads sensors, detects impact/no movement, streams compact JSON over BLE", "Vitals packets, fall events, buzzer/vibration feedback"],
                ["Flutter app", "BLE packets, user settings, guardian number, phone GPS", "Risk detection, fall countdown, movement cancellation, care suggestions, alert history", "Live dashboard, I AM OKAY reset, SOS request"],
                ["Emergency backend", "SOS payload, location, guardian target, API keys stored in .env", "Hospital lookup, message formatting, WhatsApp/push routing", "Guardian alert with Google Maps link and top medical facilities"],
            ],
            [1.1, 1.75, 2.35, 1.3],
        )
    )
    story.append(Spacer(1, 6))
    story.append(callout_flowable("Architecture flow", "ESP32-S3 Watch -> BLE JSON -> Flutter App -> 30-sec Safety Logic -> Backend -> Google Places -> WhatsApp / Push / SMS fallback", "#F2F4F7"))

    h1("4. Input Data")
    h2("4.1 Wearable Sensor Inputs")
    story.append(
        pdf_table(
            [
                ["Input", "Suggested source", "Purpose", "Current status"],
                ["Acceleration and gyro", "MPU6050 now; BMI270/ICM-42688 recommended for production", "Fall detection, movement recovery, activity context", "Firmware hooks present"],
                ["Heart rate", "MAX30102/MAX86141 optical PPG", "Sudden HR drop, cardiac load, stress support", "Optional firmware flag"],
                ["SpO2", "MAX30102/MAX86141 optical PPG", "Low oxygen risk detection", "Optional sensor calibration needed"],
                ["GSR/EDA", "Skin conductance sensor", "Stress response detection", "Analog input present"],
                ["Body temperature", "MAX30205/MAX30208 or skin temperature sensor", "Fever/heat strain context", "App meter present"],
                ["ECG amplitude", "AD8232 or compact ECG front-end", "Future ECG model input", "Analog input present; validation needed"],
                ["GPS", "Watch GPS or phone GPS fallback", "SOS location and nearby facility lookup", "Payload support present"],
                ["Battery/manual SOS", "Voltage divider/fuel gauge and watch button", "Battery warning and user-triggered alert", "Firmware inputs present"],
            ],
            [1.35, 1.85, 2.25, 1.05],
        )
    )
    h2("4.2 User, Backend, and Dataset Inputs")
    story.append(bullets(["Guardian WhatsApp number is entered by each wearable owner; it is not hard-coded.", "Backend URL, device token, and push target connect the app to the alert service.", "Google Maps and WhatsApp provider secrets stay in backend .env, not in Flutter app code.", "Doctor-suggested medicines are stored only when entered; the app should not automatically prescribe medicine."], PDF_STYLES["Body"]))
    story.append(
        pdf_table(
            [
                ["Dataset", "Signal focus", "Planned use"],
                ["MIT-BIH Arrhythmia Database", "ECG rhythm examples", "Arrhythmia research model training and validation."],
                ["Sudden Cardiac Death Holter Database", "Long ECG records around serious cardiac events", "High-risk ECG feature study; not a sole emergency rule."],
                ["Wearable Device Dataset", "BVP, EDA, temperature, motion", "Stress/activity features similar to watch input."],
                ["Wearable Exam Stress Dataset", "Stress-related wearable signals", "Stress trend validation and false-alert reduction."],
                ["Prototype watch logs", "Actual device packets", "Most important source for calibration and field performance."],
            ],
            [2.0, 1.65, 2.85],
        )
    )

    h1("5. Output Data")
    story.append(
        pdf_table(
            [
                ["Output", "Destination", "Meaning"],
                ["Live vitals", "Dashboard", "HR, SpO2, HRV, stress, GSR, ECG, body temperature, activity, movement, battery, GPS state."],
                ["AI risk detection label", "Dashboard", "User-friendly condition label such as minor stress, high temperature, low oxygen, or critical fall."],
                ["Fall countdown", "Dashboard and SOS tab", "30-second timer after fall; movement cancels it; no movement triggers SOS."],
                ["Care suggestions", "Care tab", "Exercise/home-remedy style guidance for minor criteria; serious criteria shows consult doctor."],
                ["Guardian alert", "WhatsApp/push/SMS fallback", "Reason, latest vitals, exact Google Maps link, and top nearby medical facilities."],
                ["Wearable feedback", "Watch buzzer/vibration/LED", "Immediate physical feedback during emergency countdown and connection events."],
            ],
            [1.55, 1.45, 3.5],
        )
    )

    h1("6. Emergency Logic")
    story.append(
        bullets(
            [
                "ESP32-S3 detects fall impact or sends a fall_detected event.",
                "The app starts a 30-second countdown with vibration pulses every 2 seconds and BUZZER_ON.",
                "If movement returns, the app stops the timer and sends BUZZER_OFF.",
                "If no movement remains for 30 seconds, the app fetches GPS and sends SOS.",
                "If fall plus sudden heart-rate drop is detected, immediate SOS triggers without waiting.",
            ],
            PDF_STYLES["Body"],
        )
    )

    story.append(PageBreak())
    h1("7. Hardware and Firmware Architecture")
    story.append(
        pdf_table(
            [
                ["Subsystem", "Recommended part", "Reason"],
                ["Controller", "ESP32-S3 module", "BLE support, enough processing power, common dev-board ecosystem."],
                ["Motion", "BMI270/ICM-42688 for production; MPU6050 for prototype", "Better fall and movement detection with calibrated acceleration/gyro data."],
                ["Optical vitals", "MAX86141/MAX30102", "Heart rate and SpO2 support for wrist prototypes."],
                ["Temperature", "MAX30205/MAX30208 or skin temperature sensor", "Body temperature trend input."],
                ["Stress", "GSR/EDA electrodes", "Stress trend signal when combined with HRV and movement."],
                ["ECG optional", "AD8232 or medical-grade ECG front-end", "Research input only unless clinically validated."],
                ["Location", "Phone GPS fallback plus optional watch GPS", "Best reliability because indoor GPS can fail on small wearables."],
                ["Feedback/power", "Vibration, buzzer, LED, LiPo, charger IC, fuel gauge", "Local warning plus reliable battery health monitoring."],
            ],
            [1.35, 2.15, 3.0],
        )
    )

    h1("8. Release Readiness")
    story.append(callout_flowable("Biggest risk", "The product must not miss a real emergency and must not create frequent false emergencies. This is the central safety problem to solve before public release.", "#FCE8E6"))
    story.append(
        pdf_table(
            [
                ["Risk", "Why it matters", "Control needed"],
                ["False fall alert", "Can worry guardians or waste emergency resources.", "Test with walking, running, sitting, stairs, phone drops, and sleep movement."],
                ["Missed fall", "Can delay help when the user is unconscious.", "Tune IMU thresholds using real falls/safe simulations and multiple body positions."],
                ["Loose sensor contact", "Wrist vitals can be inaccurate.", "Add signal-quality checks and show warnings when readings are unreliable."],
                ["Indoor GPS failure", "Location may be poor inside buildings.", "Use phone GPS, last-known location, Wi-Fi/network location, and accuracy labels."],
                ["Notification failure", "Guardian may not receive message.", "Use WhatsApp, push fallback, SMS composer fallback, retry logging, and delivery status."],
                ["Medical/legal exposure", "Public health products require careful claims.", "Use disclaimers, consent, privacy policy, field trials, and clinical/regulatory advice."],
            ],
            [1.35, 2.15, 3.0],
        )
    )

    h1("9. Next Milestones")
    story.append(
        bullets(
            [
                "Wait for Meta WhatsApp template approval, then test automatic WhatsApp alerts.",
                "Wire final sensors and enable firmware flags after libraries and wiring are verified.",
                "Run controlled fall/no-movement tests and tune thresholds against false alarms.",
                "Collect prototype logs and train AI models using PhysioNet plus real watch data.",
                "Prepare privacy policy, user consent, emergency-contact onboarding, and pilot-study documentation.",
            ],
            PDF_STYLES["Body"],
        )
    )

    h1("10. Local Development References")
    story.append(
        pdf_table(
            [
                ["Item", "Location"],
                ["Flutter app", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app"],
                ["Android APK", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\build\app\outputs\flutter-apk\app-debug.apk"],
                ["ESP32 firmware", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\firmware\esp32_neuroguardian_ble\esp32_neuroguardian_ble.ino"],
                ["Backend server", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\backend\emergency_server.py"],
                ["Backend secrets", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\backend\.env (do not publish)"],
                ["AI model plan", r"C:\Users\hp\Downloads\neuro gardian x\neuroguardian_app\docs\physionet_model_training_plan.md"],
            ],
            [1.65, 4.85],
        )
    )

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)


def main() -> None:
    build_docx()
    build_pdf()
    print(f"DOCX: {DOCX_OUT}")
    print(f"PDF: {PDF_OUT}")


if __name__ == "__main__":
    main()
